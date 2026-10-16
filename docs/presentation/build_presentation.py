"""Build the advisor deck and a matching PDF from saved project evidence.

Run with python-pptx, matplotlib, reportlab and pymupdf available.
No model inference, training, or experiment results are modified.
"""
from pathlib import Path
import json
import math
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
import fitz
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
STEM = 'AuditTrainer_导师汇报_2026-09-09'
FONT = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
pdfmetrics.registerFont(TTFont('WQ', FONT, subfontIndex=0))
font_manager.fontManager.addfont(FONT)
plt.rcParams.update({'font.family': 'WenQuanYi Zen Hei', 'axes.unicode_minus': False,
                     'font.size': 13, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.spines.left': False, 'axes.spines.bottom': False})

NAVY = '#17364B'
TEAL = '#07877E'
BLUE = '#4D82B8'
CORAL = '#CC684D'
GRAY = '#6B7E8B'
PALE = '#EEF3F6'
BG = '#F8FAFB'
WHITE = '#FFFFFF'
INK = '#223846'
MUTED = '#536D7C'
LIGHT_TEAL = '#E4F2EE'
LIGHT_CORAL = '#F9EDE6'

def read(p):
    return json.loads((ROOT / p).read_text())

A5 = read('runs/a5/result.json')
A6 = read('runs/a6/result.json')
P1 = read('runs/a2/mechanism_p1.json')['arms']
P2 = read('runs/a2/mechanism_p2.json')['arms']
P3 = read('runs/a2/mechanism_p3.json')['arms']
GROUPS = ['F00', 'F10', 'F01', 'F11']

def a5values(condition):
    return [A5['evaluation'][g]['conditions'][condition]['first_private_pass'] for g in GROUPS]

def barplot(name, categories, series, denominator, colors=None, ylabel=None, ymax=None):
    """Scientific charts: consistent zero baseline and full-denominator labels."""
    fig, ax = plt.subplots(figsize=(8.5, 4.3), dpi=220)
    fig.patch.set_facecolor(WHITE)
    ax.set_facecolor(WHITE)
    x = np.arange(len(categories))
    width = min(.58 / len(series), .52)
    colors = colors or [TEAL, BLUE, CORAL]
    for j, (label, values) in enumerate(series):
        off = (j - (len(series)-1)/2) * width
        bars = ax.bar(x + off, values, width*.9, label=label, color=colors[j])
        for b, v in zip(bars, values):
            ax.text(b.get_x()+b.get_width()/2, v+.022*denominator,
                    f'{v}/{denominator}', ha='center', va='bottom', fontsize=14, color=INK,
                    fontweight='bold')
    ax.set_ylim(0, ymax or denominator*1.17)
    step = 1 if denominator <= 5 else (2 if denominator <= 8 else 4)
    ax.set_yticks(np.arange(0, denominator+1, step))
    ax.set_ylabel(ylabel or f'私有通过数（分母 {denominator}）', fontsize=12, color=MUTED)
    ax.set_xticks(x, categories, fontsize=14)
    ax.tick_params(axis='both', length=0, colors=MUTED)
    ax.grid(axis='y', color='#DCE5EB', linewidth=.8)
    ax.set_axisbelow(True)
    if len(series) > 1:
        ax.legend(frameon=False, ncol=len(series), loc='upper center', bbox_to_anchor=(.5,1.16), fontsize=12)
    fig.tight_layout(pad=1.4)
    p = ASSETS / f'{name}.png'
    fig.savefig(p, facecolor=WHITE)
    plt.close(fig)
    return p

charts = {}
charts['p1'] = barplot('a2_p1', ['整体单函数', '函数级分解', '层次契约'],
    [('私有通过', [P1[k]['private_passes'] for k in ['coarse','functions','hierarchical']])],20)
charts['p2'] = barplot('a2_p2', ['事后链接', '静态计划', '逐节点契约'],
    [('私有功能通过',[P2[k]['private_passes'] for k in ['posthoc','static','hierarchical']]),
     ('来源链接完整',[P2[k]['provenance_complete'] for k in ['posthoc','static','hierarchical']])],20,
     ylabel='通过 / 完整的任务数（分母 20）')
charts['p3'] = barplot('a2_p3', ['R0 日志', 'R1 责任筛选', 'R2 候选重放'],
    [('私有修复通过',[P3[k]['private_pass'] for k in ['R0','R1','R2']])],20)
charts['i'] = barplot('a5_implementation',GROUPS,
    [('transform',a5values('transform')),('summarize',a5values('summarize'))],8)
charts['c'] = barplot('a5_composition',GROUPS,
    [('双子模块',a5values('parent_both_children')),('父节点保留汇总',a5values('parent_transform_child'))],8,
    colors=[TEAL,CORAL])
charts['d'] = barplot('a5_plan',GROUPS,
    [('计划注册',[sum(A5['evaluation'][g]['conditions']['local_plan']['accepted_plan_modes'].values()) for g in GROUPS]),
     ('固定执行器完成',a5values('local_plan'))],8,colors=[BLUE,TEAL],ylabel='成功任务数（分母 8）')
charts['free'] = barplot('a5_free',GROUPS,
    [('最终私有通过',[A5['rates'][g]['free_recursive']['successes'] for g in GROUPS])],5)

def a6plot():
    labels = ['局部计划 + 固定执行器', '固定组件合计', '双子模块组合', '父节点保留汇总（主指标）', '自由运行']
    den = [8,16,8,8,5]
    before = [3,12,7,1,1]
    after = [5,14,6,2,1]
    assert before[:1] == [A6['comparison']['local_plan']['baseline_passes']]
    assert after[1] == sum(A6['comparison'][k]['revised_passes'] for k in ['transform','summarize'])
    fig, ax = plt.subplots(figsize=(9.2,4.5),dpi=220)
    ys = np.arange(5)[::-1]
    for y, b, a, n in zip(ys,before,after,den):
        ax.plot([b/n*100,a/n*100],[y,y],color='#B4C6CE',linewidth=4,zorder=1)
        ax.scatter([b/n*100],[y],s=90,color=GRAY,label='A5 F11' if y==4 else None,zorder=3)
        ax.scatter([a/n*100],[y],s=90,color=TEAL,label='A6 F11_C' if y==4 else None,zorder=4,marker='D')
        ax.text(103,y,f'{b}/{n} → {a}/{n}',va='center',fontsize=13,color=CORAL if a<b else INK)
    ax.set_yticks(ys,labels,fontsize=13)
    ax.set_xlim(-2,126)
    ax.set_xticks([0,25,50,75,100],['0%','25%','50%','75%','100%'])
    ax.tick_params(length=0, colors=MUTED)
    ax.grid(axis='x',color='#DCE5EB',linewidth=.8)
    ax.set_axisbelow(True)
    ax.legend(frameon=False,ncol=2,loc='upper center',bbox_to_anchor=(.55,1.16))
    ax.set_xlabel('私有通过率；右侧保留实际分子 / 分母',fontsize=12,color=MUTED)
    fig.tight_layout(pad=1.4)
    p=ASSETS/'a6_comparison.png'
    fig.savefig(p,facecolor=WHITE)
    plt.close(fig)
    return p
charts['a6']=a6plot()

def wrap(text, size, width):
    lines=[]
    for para in text.split('\n'):
        if not para:
            lines.append('')
            continue
        line=''
        for c in para:
            if pdfmetrics.stringWidth(line+c,'WQ',size)>width and line:
                lines.append(line.rstrip())
                line=c.lstrip()
            else:
                line+=c
        lines.append(line.rstrip())
    return lines

class Deck:
    def __init__(self):
        self.prs=Presentation()
        self.prs.slide_width=Pt(960)
        self.prs.slide_height=Pt(540)
        self.pdf=canvas.Canvas(str(OUT/f'{STEM}.pdf'),pagesize=(960,540))
        self.pdf.setTitle('AuditTrainer 阶段汇报：框架、实验与结果解释')
        self.pdf.setAuthor('AuditTrainer')
        self.slides=[]
        self.bounds=[]

    def rect(self,x,y,w,h,fill=WHITE,line=None,radius=0):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                      Pt(x),Pt(y),Pt(w),Pt(h))
        sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill.lstrip('#'))
        if line:
            sh.line.color.rgb=RGBColor.from_string(line.lstrip('#')); sh.line.width=Pt(.7)
        else: sh.line.fill.background()
        if radius:
            sh.adjustments[0]=.05
        self.pdf.setFillColor(HexColor(fill))
        self.pdf.setStrokeColor(HexColor(line or fill))
        self.pdf.roundRect(x,540-y-h,w,h,radius,stroke=bool(line),fill=1)
        return sh

    def text(self,x,y,w,h,text,size=18,color=INK,bold=False,align='left',leading=1.28):
        lines=wrap(text,size,w-3)
        needed=len(lines)*size*leading
        if needed>h+1:
            raise ValueError(f'Slide {len(self.slides)} text overflow: need {needed:.1f}, have {h}: {text}')
        self.bounds.append((len(self.slides),x,y,w,h,text))
        box=self.slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
        tf=box.text_frame
        tf.clear(); tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(size*leading)
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            r=p.add_run(); r.text=line
            r.font.name='Microsoft YaHei'; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
            for tag in ['a:ea','a:cs']:
                el=OxmlElement(tag); el.set('typeface','Microsoft YaHei'); r._r.get_or_add_rPr().append(el)
            tw=pdfmetrics.stringWidth(line,'WQ',size)
            xx=x if align=='left' else (x+(w-tw)/2 if align=='center' else x+w-tw)
            t=self.pdf.beginText(xx,540-y-size-i*size*leading)
            t.setFont('WQ',size); t.setFillColor(HexColor(color)); t.setStrokeColor(HexColor(color))
            if bold: t.setTextRenderMode(2); self.pdf.setLineWidth(.16)
            t.textOut(line); self.pdf.drawText(t)
        return box

    def image(self,x,y,w,h,path):
        self.slide.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(w),height=Pt(h))
        self.pdf.drawImage(str(path),x,540-y-h,width=w,height=h,mask='auto')

    def start(self,section,title,subtitle,notes,source,takeaway=None,dark=False):
        if self.slides: self.pdf.showPage()
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slides.append({'section':section,'title':title,'notes':notes,'source':source})
        self.slide.notes_slide.notes_text_frame.text=(f'{title}\n\n{notes}\n\n证据来源：\n'+'\n'.join(str(ROOT/p) for p in source))
        self.rect(0,0,960,540,NAVY if dark else BG)
        if not dark:
            self.rect(0,0,8,540,TEAL)
            self.text(42,22,850,20,section.upper(),11,TEAL,True)
            self.text(42,50,880,43,title,29,NAVY,True)
            self.text(42,99,878,26,subtitle,13,MUTED)
            if takeaway:
                self.rect(42,471,878,38,LIGHT_TEAL,radius=6)
                self.text(56,480,850,23,takeaway,15,TEAL,True)
        self.text(42,520,830,15,'来源：'+'；'.join(source),8.7,'#B6C9D4' if dark else GRAY)
        self.text(900,519,28,17,f'{len(self.slides):02d}',10,'#B6C9D4' if dark else GRAY,align='right')

    def card(self,x,y,w,h,tag,title,body,color=TEAL):
        self.rect(x,y,w,h,WHITE,radius=6)
        self.rect(x,y,4,h,color)
        self.text(x+18,y+15,w-36,22,tag,11,color,True)
        self.text(x+18,y+45,w-36,55,title,23,NAVY,True)
        self.text(x+18,y+111,w-36,h-119,body,16,MUTED)

    def bullets(self,x,y,w,items,size=18,gap=16,color=INK):
        for item in items:
            lines=wrap(item,size,w-24)
            h=len(lines)*size*1.28+2
            self.rect(x,y+8,5,5,TEAL)
            self.text(x+17,y,w-17,h,item,size,color)
            y+=h+gap

    def table(self,x,y,widths,headers,rows,rowh=42,size=16,colors=None):
        total=sum(widths)
        self.rect(x,y,total,rowh,NAVY)
        xx=x
        for head,w in zip(headers,widths):
            self.text(xx+10,y+10,w-20,rowh-12,str(head),size-1,WHITE,True)
            xx+=w
        for i,row in enumerate(rows):
            yy=y+(i+1)*rowh
            self.rect(x,yy,total,rowh,WHITE if i%2==0 else PALE)
            xx=x
            for j,(cell,w) in enumerate(zip(row,widths)):
                cc=(colors or {}).get((i,j),INK)
                self.text(xx+10,yy+10,w-20,rowh-12,str(cell),size,cc,bold=j>0 and bool(re.search(r'\d/\d',str(cell))))
                xx+=w

    def save(self):
        self.prs.core_properties.title='AuditTrainer：框架、实验与结果解释'
        self.prs.core_properties.subject='截至 2026-09-09；20 页正文 + 5 页备份'
        self.prs.core_properties.author='AuditTrainer'
        self.prs.core_properties.keywords='责任分解,契约履约,有界修复,导师汇报'
        self.prs.save(OUT/f'{STEM}.pptx')
        self.pdf.save()
        md=['# AuditTrainer PPT 逐页讲解稿','',
            '对应 25 页 PPT：第 1–20 页为正文，第 21–25 页为备份。建议正文讲解约 20 分钟；短汇报可略讲第 5、9、10、12 页。',
            '', '每页讲解已写入 PPT 演讲者备注。图中均保留原始分子/分母；无显著性或最终泛化声明。','']
        for i,s in enumerate(self.slides,1):
            md += [f'## 第 {i:02d} 页｜{s["title"]}','',s['notes'],'',
                   '证据：'+ '、'.join(f'[{p}]({ROOT/p})' for p in s['source']),'']
        (OUT/f'{STEM}_逐页讲解稿.md').write_text('\n'.join(md))
        for page,x,y,w,h,txt in self.bounds:
            assert x>=0 and y>=0 and x+w<=961 and y+h<=541,(page,txt)

D=Deck()

# 01 — Cover
D.start('阶段汇报','AuditTrainer','',
'''开场：目前已完成从运行框架到核心训练与开发评测的闭环，阶段推进到 A6。研究围绕责任分解、契约履行与组合、有界修复展开。
最清楚的结果来自 A5：F01 相对普通功能 SFT 对照 F00，固定组件通过数从 0/16 提高到 13/16；双子模块组合从 1/8 提高到 8/8。随后改变父子职责分配，暴露出父节点重新承担汇总职责时的不足。
今天汇报重点是实验为什么这样设计、这些结果支持什么判断，以及下一步如何补齐结构对修复的价值证据。所有结果截至 2026-09-09，属于单 seed 开发比较。''',
['STATUS.md','reports/A5.md','reports/A6.md'],dark=True)
D.text(48,40,810,30,'研究进展  /  2026.09.09',14,'#87CFC5',True)
D.text(48,114,865,90,'AuditTrainer',57,WHITE,True)
D.text(48,213,865,62,'责任分解 · 契约履约 · 有界修复',30,WHITE,True)
D.text(48,281,840,58,'已完成框架闭环与核心训练对照\n重点：实验设计、结果解释与下一步证据',20,'#C3D4DF')
for x,num,label in [(48,'A0–A6','当前完成阶段'),(347,'5 组','正式训练模型'),(646,'225 次','A5/A6 开发运行')]:
    D.rect(x,381,268,101,'#24465B',radius=6)
    D.text(x+17,392,234,42,num,30,'#8EDBD0',True)
    D.text(x+17,444,234,23,label,14,WHITE)

# 02 — Main results
D.start('01 / 研究进展','目前最清楚的成果：局部履约与组合可被训练改善',
        '主结果来自 A5；自由递归和修复收益仍需分别验证。',
'''目的：先让导师建立对当前进展的整体判断。这里把工程完成、科学正向结果和能力瓶颈分开表达。
结果：50 个 TRAIN 根任务形成 2,240 条训练动作；四个 A5 模型加一个 A6 修订模型均正式训练 140 步并回接运行时。F01 对比 F00 的组件和双子模块结果是当前最明确的正向证据。
解释：履约训练让模型更稳定地生成符合局部接口的代码，并把固定组件串联起来。改变父节点职责后成功下降，说明常见组合行为尚未形成稳定的职责适应能力。
边界：0/16→13/16 是两种训练组之间的开发对照，F00 也经过 SFT；不能写成公开底座直接提升到 81.25%，也不是多 seed 显著性结论。''',
        ['reports/A4.md','runs/a5/result.json','runs/a6/result.json'],
        '核心判断：固定契约下的履约收益明确；职责变化后的适应仍是瓶颈。')
D.card(42,145,282,305,'已完成','完整研究闭环','运行 → 数据 → 训练 → 评测\n\n50 个 TRAIN 根任务\n2,240 条动作\n5 组最终 adapter')
D.card(340,145,282,305,'已观察到','0/16 → 13/16','F00 → F01 固定组件\n\n双子模块组合\n1/8 → 8/8\n\nI 监督产生局部收益')
D.card(638,145,282,305,'关键瓶颈','8/8 → 0/8','F01 双子模块 → 单子模块\n\n父节点重新承担汇总后\n常见组合能力未能保持\n\n下一步连接到修复价值',CORAL)

# 03 — Research question
D.start('02 / 研究问题','责任是否真正改变了模型的实现行为？',
        '说明性示意：根任务相同，父子职责分配不同；与 A5 的两类父组合条件对应。',
'''目的：解释为什么要单独测“责任”，以及为什么职责重分配是关键实验。
设计：双子模块条件中，transform 负责变换，summarize 负责汇总，父节点负责解析、数据传递与格式化。单子模块条件保留同一个 transform 子契约，把汇总职责合法地留给父节点；根任务与要求的最终结果不变。
解释：父节点需要根据当前 obligation map 和 retained responsibilities 改变代码。若只是复现训练中的传递模式，就可能直接返回中间列表、遗漏汇总或沿用错误输出包装。
另一个区别是父节点可能调用 child 后丢弃返回值，自己重算。调用记录只能证明执行过，不能单独证明遵守了责任分配。A4 用受控对照明确展示这种差异。
本页为机制示意，不是额外实验结果。''',
        ['reports/A4.md','reports/A5.md','src/audittrainer/schema.py'],
        '评测对象：局部算法正确 + 按当前契约分工 + 正确使用子产物。')
D.rect(42,145,430,302,WHITE,radius=6); D.rect(490,145,430,302,WHITE,radius=6)
D.text(62,164,390,35,'双子模块：委托变换与汇总',21,NAVY,True)
D.text(510,164,390,35,'单子模块：父节点保留汇总',21,NAVY,True)
for x in [62,510]:
    D.rect(x,221,390,48,PALE,radius=5); D.text(x+14,232,362,28,'父节点：解析输入',18,NAVY)
D.rect(62,288,180,48,LIGHT_TEAL,radius=5); D.text(73,299,158,28,'transform 子模块',16,TEAL,True)
D.text(249,297,25,32,'→',20,GRAY)
D.rect(280,288,172,48,LIGHT_TEAL,radius=5); D.text(289,299,157,28,'summarize 子模块',15,TEAL,True)
D.rect(510,288,390,48,LIGHT_TEAL,radius=5); D.text(524,299,362,28,'transform 子模块',18,TEAL,True)
D.rect(62,355,390,66,PALE,radius=5)
D.text(76,371,362,32,'父节点：格式化最终结果',18,NAVY)
D.rect(510,355,390,66,LIGHT_CORAL,radius=5)
D.text(524,367,362,47,'父节点：自己汇总\n再按当前输出契约格式化',17,CORAL,True)

# 04 — Framework
D.start('03 / 已完成框架','一个共享策略，贯通计划、执行、证据与修复',
        '同一 Qwen3-8B 服务全部递归节点；运行时维护契约、产物版本和根任务共享预算。',
'''目的：说明已经实现的系统如何支持三条研究主线。
流程：模型先输出 ATOMIC 或 SPLIT 计划，约定子契约、义务归属和父节点保留责任。子节点递归执行，父节点通过 consume 选择精确版本，再由模型自己编写组合代码。可见测试返回与当前契约和产物绑定的证据。
修复：R1 责任筛选提供候选，模型输出 localize 和 repair；ticket 定义真实可编辑边界，边界外源码冻结。支持代码补丁和内部计划子树重建，以可见检查全过且无已知回归为接受条件，否则回滚。
主配置：深度从根 0 到最大 2，最多 7 个累计节点、每节点 3 个 child，根共享 64 次策略调用、32 次验证/重放、3 次补丁。框架支持递归与修复，并不代表当前训练模型已经主动使用这些能力。
私有评分只在程序快照选定、模型决策结束后执行。''',
        ['configs/frozen_mechanism.yaml','reports/F3.md','reports/F4.md'],
        '已经具备实验执行能力；模型是否学会分解和修复，由相应实验单独判断。')
for x,tag,title,body in [(42,'D / 分解','契约 → 责任计划','ATOMIC / SPLIT\n子契约与义务映射\n父节点保留职责'),
                         (340,'I + C / 履约组合','实现 → 消费 → 验证','每节点一个源码模块\n版本 / hash / alias 绑定\n父组合代码由模型生成'),
                         (638,'R / 有界修复','定位 → 补丁 → 重放','R1 责任筛选\nticket 指定实际修改范围\n冻结补集，接受或回滚')]:
    D.card(x,147,282,246,tag,title,body)
D.rect(42,410,878,39,PALE,radius=6)
D.text(56,419,850,25,'可见证据支持在线决策    |    首个 / 最终快照 → 离线私有评分',17,NAVY,True,align='center')

# 05 — Experiment map
D.start('04 / 实验地图','每个阶段回答一个明确的问题',
        '主线已经推进到 A6；A7 是下一步实验，尚无结果。',
'''目的：把后续结果放回研究逻辑，避免仅按工程阶段流水介绍。
A1 验证运行、训练和回接能否真实执行；A2 用 20 题比较责任粒度、承诺时机和修复候选机制；A3 根据 A2 冻结唯一主方案。
A4 构造 50 个训练任务及可追踪示范、对照，解决监督来源问题；A5 在同样训练预算下做四组两因素比较，并拆开测局部计划、组件、组合和自由运行；A6 针对父职责覆盖做一次有限干预。
下一步 A7 固定失败和修复器，检验责任结构的下游价值。160、180、45 都是条件运行次数，不是独立新任务数，A5/A6 的开发任务重复使用。''',
        ['STATUS.md','reports/A2.md','reports/A5.md','reports/A6.md'],
        '实验链条：机制诊断 → 可审查监督 → 能力隔离 → 一次修订 → 修复价值。')
D.table(42,145,[89,272,315,202],['阶段','要回答的问题','实际设置','规模 / 状态'],[
    ['A1/A3','框架能否支撑研究？','运行、训练、回接；冻结主机制','已完成'],
    ['A2','选择什么机制？','五种生成 + 三种修复条件','20 题 / 160 次'],
    ['A4','责任监督从哪里来？','三种计划、独立执行、六类对照','50 题 / 2,240 动作'],
    ['A5','训练改善哪种能力？','D × I 四组；隔离 + 自由运行','4 × 140 步 / 180 次'],
    ['A6','补父职责覆盖有用吗？','仅替换父组合槽位；原协议复评','140 步 / 45 次'],
    ['A7','结构能否改善修复？','固定失败、固定修复器两轨','下一阶段 / 未完成'],
],rowh=43,size=14)

# 06 — P1
D.start('05 / A2 · P1 责任粒度','局部组件可通过，完整组合仍会失败',
        '同一未微调模型、同一批 20 题；比较整体单函数、函数级分解和层次契约。',
'''目的：比较责任表示粒度，判断函数/模块级结构能否作为后续学习载体。
设计：coarse 在一个 solve 函数里完成任务；functions 提供函数接口但不要求父义务映射；hierarchical 提供完整子契约、义务分配与显式消费。模型负责代码、非平凡计划和消费，共用固定执行控制。
结果：三组私有成功分别 8/20、2/20、1/20。functions 与 hierarchical 的局部检查都达到 29/40，但分别有 7/20 和 9/20 个任务在两个 child 都通过后仍未形成根程序。
解释：完整运行瓶颈集中在消费、组合和协议履行，不能把低根成功率简单归为所有子模块算法都不会。此结果支持优先补履约与消费数据。
边界：functions 仍生成了不符合条件的义务映射，严格遵守率为 0/20；hierarchical 为 20/20。该实验改变表示、信息和交互结构，不能估计单字段因果效应，也没有显示层次契约的零样本生成优势。''',
        ['reports/A2.md','runs/a2/mechanism_p1.json'],
        'A2 的直接作用是定位训练缺口：履约、消费与父组合需要针对性监督。')
D.image(42,153,552,280,charts['p1'])
D.rect(612,145,308,302,WHITE,radius=6)
D.text(631,160,270,42,'29/40',31,TEAL,True)
D.text(631,206,270,48,'函数级 / 层次契约\n均有这些局部检查通过',16,MUTED)
D.text(631,269,270,56,'但 7/20、9/20 题\n双 child 通过后仍无根程序',18,NAVY,True)
D.text(631,351,270,74,'条件遵守：\nfunctions 0/20\nhierarchical 20/20',16,CORAL)

# 07 — P2
D.start('06 / A2 · P2 承诺时机','来源链接完整，不能代替功能或责任正确',
        '比较先代码后链接、先静态计划、逐节点契约承诺；每组同样 20 题。',
'''目的：检验生成后补链接与生成期承诺之间的差异，避免把“有责任标注”直接当作履约。
设计：posthoc 先冻结完整代码，再给实际函数补来源链接；static 先声明 helper 和义务关系，再生成单模块代码；hierarchical 逐节点提交契约、实现和消费。hierarchical 与 P1 共享同一批轨迹，统计总运行数时不重复计数。
结果：功能通过依次为 9/20、6/20、1/20；来源链接完整为 1/20、20/20、2/20。静态计划即使链接完整，功能通过也只有 6/20。
解释：链接完整仅说明义务指向实际源码对象，并不验证算法正确、分工语义或父模块使用行为。当前数据未支持逐节点承诺的零样本生成优势；保留这一表示用于后续训练检验。
边界：static 对预定 helper/映射的严格遵守为 0/20，因此不是完全实施的纯静态机制对照。不能从柱状图得出责任标签无用，也不能把静态链接 20/20 称为责任正确 20/20。''',
        ['reports/A2.md','runs/a2/mechanism_p2.json'],
        '来源链接与语义履约分别测量，训练标签需要独立执行和责任证据。')
D.image(42,160,590,285,charts['p2'])
D.bullets(651,165,254,[
    '静态计划：链接 20/20，功能仅 6/20。',
    '链接证明“指向谁”；不能证明“做对且按约做”。',
    'static 条件遵守 0/20，保留解释边界。',
],size=17,gap=23)

# 08 — P3
D.start('07 / A2 · P3 受控修复','责任筛选在共同故障上取得较好的修复结果',
        '同一批 20 个受控故障、同一补丁模型和总预算；私有评分不参与补丁选择。',
'''目的：在共同失败起点下比较日志定位、责任筛选和有限候选重放。
设计：R0 读取日志与源码；R1 增加责任索引和由可见失败筛得的候选；R2 增加无修改重放、独立候选比较与 ticket 冻结。每题最多 3 个补丁、24 次验证、10 次模型调用和 65,536 输入加输出 token。
结果：R0/R1/R2 首选注入位置命中为 14/20、20/20、20/20；私有修复为 17/20、19/20、18/20。R1 相对 R0 多成功 2 题，补丁尝试从 31 降到 26；在线验证从 93 降到 78。R2 在线验证 138 次，没有额外私有成功。
解释：R1 是本次有限机制比较中更合适的定位起点；候选轮换重放在当前任务与策略下没有显示新增价值。责任索引增加输入长度，因此不能声称所有成本都下降。
关键边界：R0/R1 虽声明首选 owner，实际授权三个模块；R2 每 ticket 只授权一个。A2 的 R1 结果不能直接当作 A3 冻结后的“R1 + 强制冻结补集”组合性能。R2 有一次越界补丁被拒绝；三组均无已观测用例回归。''',
        ['reports/A2.md','runs/a2/mechanism_p3.json'],
        '开发选择：R1 作为定位起点；“R1 + 强制冻结”的完整价值继续由 A7 检验。')
D.image(42,149,518,268,charts['p3'])
D.table(579,151,[144,61,61,65],['指标','R0','R1','R2'],[
    ['位置命中','14/20','20/20','20/20'],
    ['补丁尝试','31','26','27'],
    ['在线验证','93','78','138'],
    ['实际授权模块','3','3','1'],
],rowh=43,size=13)
D.text(579,384,330,65,'三组均无已观测回归。\nR1 输入更长，token 成本未全面降低。',15,MUTED)

# 09 — A4 data
D.start('08 / A4 · 数据构造','把责任计划、局部实现和父组合变成训练样本',
        '50 个训练根任务，10 个任务族 × 每族 5 个参数配置；同族保留在 TRAIN。',
'''目的：为 D、I/C 提供可追踪、可审查的一步动作监督。
设计：每个 root 在同一父状态下执行 ATOMIC、两个 child 流水组合、单个 fused child 三种合法计划，使用同一个参考执行器。参考执行器是审查过的程序构造，不是真实模型采样。候选代码经隔离执行，以另一份算法实现的 oracle 检查。TRAIN 私有检查只在终止后用于标签。
结果：150 条参考示范全部通过；另外每族固定一个 root 采集 K=1 的底座自然轨迹，共 10 条，均提出超出已显示 child_slots=3 的四 child 计划，在反馈后弃权，没有完整根程序。自然正例不足，所以正式训练主体来自参考示范。
导出：2,240 条动作包括 D_PLAN=460、D_IMPL=320、D_COMPOSE=400、D_CONTROL=1,060。责任支持 1,180 条，协议监督 1,060 条；控制动作正确不意味着代码责任已被证明。
解释：本阶段证明数据和训练链路有可信来源；150/150 不能解释为底座模型达到 100% 成功率。50 个参数化任务也不是 50 个独立任务族。''',
        ['reports/A4.md','runs/a4/result.json','data/sft/train.manifest.json'],
        '2,240 条完整动作来自真实执行；参考示范与自然模型轨迹明确区分。')
D.card(42,145,282,303,'同父状态 / 同执行器','50 × 3 = 150','整体实现 ATOMIC\n两个 child 流水组合\n一个融合 child\n\n150 条参考示范全部通过')
D.table(340,145,[173,100],['数据桶','动作数'],[
    ['计划 D_PLAN','460'],['局部 D_IMPL','320'],['组合 D_COMPOSE','400'],['控制 D_CONTROL','1,060'],['合计','2,240'],
],rowh=45,size=15)
D.rect(634,145,286,303,WHITE,radius=6)
D.text(652,162,250,37,'自然轨迹：0/10 完整程序',19,CORAL,True)
D.text(652,217,248,94,'全部因超 child 槽位计划\n在反馈后弃权。\n正式数据以参考示范为主。',17,MUTED)
D.text(652,337,248,87,'1,180 条责任支持\n1,060 条协议监督\nUNKNOWN 不作违约负例',17,NAVY,True)

# 10 — A4 controls
D.start('09 / A4 · 受控对照','功能匹配的责任违规，检验了独立的责任维度',
        '六类单因素变换各 10 对，共 60 对；保留原始、变换、执行与审查证据。',
'''目的：检验责任标签是否能区分功能错误、责任违规和协议违规，尤其是输出正确但分工被破坏的情况。
设计与结果：遗漏输出、错误参数传递形成 20 对功能失败；算法迁移到 parent、调用 child 后丢弃结果并重算形成 20 对功能匹配责任违规；选择真实旧 artifact、扩大 child owns 形成 20 对协议/契约违反。每类各 10 对，全部真实执行。
解释：其中最重要的对照是“调用后弃用”：called=true 不足以支持责任正确。独立源码审查显示 parent 重算了本来委托给 child 的算法，即使根输出匹配也属于责任问题。这为后续责任监督提供了具体判据。
质量：用户报告人工复核通过，Codex 二次复核覆盖真实计划和全部负例；没有把总体复核反馈回填成逐事件真人标注。
边界：这些是受控构造，不是对任意程序语义的通用自动证明；不同执行 observation 不完全相同，preference_ready=true 为 0，所以不能直接作为同输入 DPO 偏好对。''',
        ['reports/A4.md','runs/a4/contrasts.jsonl'],
        '关键证据：相同根输出与实际调用，并不足以证明按既定责任分工完成任务。')
D.table(42,148,[230,272,130,246],['变换','要检查的区别','数量','结果类别'],[
    ['遗漏必需输出','输出契约未完成','10 对','功能失败'],
    ['错误参数传递','未传递 prepare 结果','10 对','功能失败'],
    ['把子算法迁移到父节点','功能相同，责任归属变化','10 对','功能匹配责任违规'],
    ['调用后丢弃并重算','调用存在，返回值未使用','10 对','功能匹配责任违规'],
    ['使用真实旧产物版本','选择过期版本','10 对','协议 / 契约违反'],
    ['扩大 child 的 owns','声明超出原义务分配','10 对','协议 / 契约违反'],
],rowh=41,size=14)

# 11 — A5 design
D.start('10 / A5 · 训练设计','四组训练分开检验 D 与 I 的针对性监督',
        'F00 也是普通功能 SFT 对照；关闭因素时，使用同 root 新执行的普通示范替换。',
'''目的：分离责任分解示范 D 与局部履约/父组合示范 I 的训练影响。
设计：F00 关闭 D 和 I，F10 仅开启 D，F01 仅开启 I，F11 同时开启。关闭 D 的槽位替换为基础 ATOMIC 计划；关闭 I 的 write_code 槽位用同 root 的普通整题功能示范替换。consume 与控制动作完全共享，不加入故意违反责任的坏目标。
匹配：四组使用同一公开 Qwen3-8B revision、seed 42、50 个 TRAIN root、2,240 个共同槽位和 140 次更新；均从底座新训，取最终 adapter，不继承 pilot 或挑中间 checkpoint。
解释：该设计检验训练覆盖因素的贡献。它不是仅改标签的消融：不同组的输入/目标 token、去重数据池、重复曝光次数不同。这些实际剂量差异在备份页给出。
训练：NF4/BF16，all-linear LoRA r16/alpha32/dropout0.05，约 43.65M 可训练参数，完整 action + EOS 损失。训练 loss 只验证优化过程，不据组间 loss 比较能力。''',
        ['reports/A5.md','runs/a5/data_comparison.json','configs/a5/F11.yaml'],
        '匹配 root、采样槽位和更新次数；输入 / 目标 token 与不同样本曝光并非完全匹配。')
for x,y,g,desc,c in [(42,151,'F00','D 关闭  /  I 关闭\n普通功能示范对照',GRAY),
                      (340,151,'F10','D 开启  /  I 关闭\n分解监督',BLUE),
                      (42,302,'F01','D 关闭  /  I 开启\n履约与组合监督',TEAL),
                      (340,302,'F11','D 开启  /  I 开启\n联合监督',TEAL)]:
    D.rect(x,y,282,135,WHITE,radius=6); D.rect(x,y,4,135,c)
    D.text(x+18,y+13,246,42,g,27,c,True)
    D.text(x+18,y+62,246,60,desc,17,MUTED)
D.bullets(651,158,255,['同一公开 Qwen3-8B','每组 50 个 TRAIN root','2,240 次样本曝光','140 次参数更新','共同 seed 42'],size=18,gap=17)

# 12 — A5 evaluator
D.start('11 / A5 · 能力隔离','把“不会分解”与“不会实现 / 组合”分开测',
        '每组：8 个开发 root × 5 条件 = 40 次隔离；另有 5 个工程任务自由运行。',
'''目的：整条自由轨迹失败可能由计划格式、子算法、消费或父组合造成，必须通过固定输入条件定位能力来源。
五个条件：① local_plan 由被评模型生成根计划，再交由固定公开底座 E_ref 实现；② transform 固定局部契约，只测组件实现；③ summarize 同理；④ parent_both_children 固定两个只读子产物，只测 parent 消费和组合；⑤ parent_transform_child 只提供 transform 子产物，parent 需要自行履行 summarize 责任。
自由运行：全部决策由被评模型作出，使用 5 个 DEV_ENGINEERING root，即两个自建任务与三个 APPS 开发题。
解释：局部计划+E_ref 的成功率仍包含固定执行器的执行能力，并不是纯分解正确率。固定组件的 16 个条件是 8 个 root 各两种组件，不是 16 个独立任务。
注意 A5 的 E_ref 是公开未微调模型；A4 的参考示范执行器是已审查的程序构造，两者不能混称。所有私有分数在产物选择后计算，失败和无程序都保留分母。''',
        ['configs/a5_evaluation.yaml','reports/A5.md'],
        '隔离条件固定执行器或子产物，帮助判断收益发生在哪一项能力。')
D.table(42,146,[58,250,360,210],['条件','被测能力','固定部分','每组次数'],[
    ['①','局部计划','公开底座 E_ref 执行后续','8'],
    ['②','transform 履约','局部契约与接口','8'],
    ['③','summarize 履约','局部契约与接口','8'],
    ['④','双子模块消费 / 组合','两个 child 的契约与只读产物','8'],
    ['⑤','父节点保留汇总','transform child 契约与只读产物','8'],
    ['自由','完整策略行为','同一任务集合与总预算','5'],
],rowh=43,size=15)

# 13 — A5 I
D.start('12 / A5 · 固定组件履约','I 监督带来最清楚的局部实现收益',
        '每组固定同样 8 个 transform 和 8 个 summarize 契约；分母包含无程序与协议失败。',
'''目的：在责任契约已给定时，单独测模型能否写出符合接口和语义要求的局部实现。
结果：F00 的 transform/summarize 为 0/8、0/8；F10 为 3/8、0/8；F01 为 7/8、6/8；F11 为 6/8、6/8。两种组件合计依次为 0/16、3/16、13/16、12/16。
解释：开启 I 的两组都出现明显改善，支持履约监督对固定契约实现有作用。F00 的 transform 有 7 个任务因为公开函数接口不匹配而无程序，说明提升涉及按契约生成接口，不仅是同接口下的算法准确率。
描述性因素效应：I 的平均差为 ((13/16-0/16)+(12/16-3/16))/2=68.75 个百分点。此处只描述当前 seed 和开发条件，不报告统计显著性。
剩余错误：F11 的 prefix_span 曾调用未定义 helper，局部算法和接口错误并未完全消除。F11 没有超过 F01，不能声称 D 与 I 联合一定更好。''',
        ['reports/A5.md','runs/a5/result.json'],
        '结果支持固定契约履约的训练收益；联合 F11 尚未优于仅开启 I 的 F01。')
D.image(42,153,605,292,charts['i'])
D.rect(668,148,252,302,WHITE,radius=6)
D.text(687,168,215,25,'两类组件合计',16,MUTED,True)
for y,g,v,c in [(214,'F00','0/16',GRAY),(261,'F10','3/16',BLUE),(308,'F01','13/16',TEAL),(355,'F11','12/16',TEAL)]:
    D.text(687,y,88,31,g,20,c,True); D.text(774,y,123,34,v,24,c,True)
D.text(687,407,216,26,'I 平均差：+68.75 pp',16,TEAL,True)

# 14 — A5 C
D.start('13 / A5 · 父组合与责任重分配','常见组合可以成功，职责变化后仍不可靠',
        '同样的根任务：双子模块时父节点只传递结果；单子模块时父节点另需完成汇总。',
'''目的：区分常见父组合模式与随当前契约重新承担责任的能力。
结果：F00/F10/F01/F11 双子模块功能通过为 1/8、1/8、8/8、7/8；只有 transform 子模块、父节点保留汇总时为 0/8、0/8、0/8、1/8。F01 的 8/8 对 0/8、F11 的 7/8 对 1/8 构成最明显的落差。
责任证据：F01/F11 的 15 个双子模块成功程序均有直接消费、传递返回值和只做解析/格式化的源码证据。F00/F10 虽各有一个功能通过，未满足该保守责任判据，不能把它们也算责任成功。UNKNOWN 不等于违规。
解释：履约示范帮助模型组合固定 child，但当前训练没有充分覆盖父节点保留 aggregate 的状态。单子模块失败表现为直接返回 transform 的列表、错误汇总，以及沿用训练题中的对象包装而不是当前所需 JSON 数值。
这支持一个具体覆盖缺口，尚不能仅由这组结果判定所有错误都来自模板记忆。A6 用一次数据干预检验该解释。''',
        ['reports/A5.md','runs/a5/result.json'],
        '最值得强调的发现：组合成功尚未稳定迁移到父节点重新承担职责的状态。')
D.image(42,155,610,289,charts['c'])
D.rect(670,148,250,302,WHITE,radius=6)
D.text(689,168,213,38,'F01：8/8 → 0/8',22,CORAL,True)
D.text(689,215,213,42,'F11：7/8 → 1/8',22,CORAL,True)
D.bullets(689,282,211,['遗漏父汇总','汇总算法错误','输出包装不符当前契约'],size=17,gap=14)

# 15 — A5 D
D.start('14 / A5 · 局部计划','当前没有观察到主动递归分解的学习收益',
        '被评模型产生 plan，固定公开底座 E_ref 完成后续实现；拆开报告计划注册与最终成功。',
'''目的：测分解训练是否改善计划形成，并用固定执行器减少后续执行模型差异。
结果：F00/F10/F01 的 8 个计划全部被接受，F11 只有 4 个被接受；之后由固定 E_ref 执行，私有成功分别为 5/8、5/8、5/8、3/8。所有被接受的计划均为 ATOMIC，没有主动形成 SPLIT。
失败解释：F11 四次失败发生在注册前，例如 obligation_map 的值写成字符串 self，而 schema 要求数组 [self]，一次反馈重试后仍未修正。有效计划之后 E_ref 的失败数分别为 3、3、3、1，需要单列，不能统称分解语义失败。
结论：D 监督没有在该隔离实验中形成可见递归分解收益，联合 F11 还存在格式稳定性问题。局部计划通过率的变化不能直接代表语义分解提升。
后续 A6 中注册失败消失，但计划仍全为 ATOMIC；因此 A6 局部计划成绩改善也不等于模型学会分解。''',
        ['reports/A5.md','runs/a5/result.json'],
        '先区分格式失败、有效原子计划后的执行失败，再讨论分解能力。')
D.image(42,155,604,289,charts['d'])
D.rect(664,148,256,302,WHITE,radius=6)
D.text(683,166,220,38,'全部为 ATOMIC',22,CORAL,True)
D.text(683,224,219,102,'F11：4 个计划未注册\n\nobligation_map 值类型错误\n字符串 vs 数组',17,MUTED)
D.text(683,353,219,75,'计划被接受后，\nE_ref 仍可能执行失败。',18,NAVY,True)

# 16 — A5 free
D.start('15 / A5 · 自由运行','隔离条件的提升尚未形成端到端递归收益',
        '5 个固定 DEV_ENGINEERING root：两个自建任务 + 三个 APPS 开发题。',
'''目的：检查各模型在全部决策自由的情况下，局部收益是否转化为预算内整题成功。
结果：F00/F10/F01/F11 首个程序私有成功与最终成功均分别为 1/5、1/5、2/5、1/5。四组共 20 个运行全部保留，评分没有基础设施 ERROR。
行为：所有自由轨迹只形成 ATOMIC，没有实际修复尝试，所以 FirstProgramSuccess 与 BudgetedSuccess 相同。这说明模型没有主动调用框架已具备的递归和修复机制，不能据此评价修复是否有效。
解释：F01 比其他组多通过一个任务是描述性结果，只有 5 个开发 root，不应报告为稳定端到端优势。训练对固定组件的收益尚未充分转化成自由策略能力。
范围：三个 APPS 题来自已消费 train 开发记录，不代表整个 APPS 基准；其离线评分也不保证对预训练未见。早期 F9 的 base 2/5 使用不同阶段设置，不能拿来充当这轮严格匹配的额外基线。''',
        ['reports/A5.md','configs/a5_evaluation.yaml','reports/F9.md'],
        '首个成功 = 最终成功，是因为本轮没有实际修复尝试；不能解释为修复无效。')
D.image(42,152,586,290,charts['free'])
D.card(649,148,271,302,'完整策略行为','均为 ATOMIC','全部 20 个运行保留\n无实际修复尝试\n责任状态均 UNKNOWN\n\nF01 多通过 1 个任务\n仍是小规模开发结果',CORAL)

# 17 — A6 design
D.start('16 / A6 · 一次数据修订','只补父节点保留汇总的示范覆盖',
        '沿用 A5 诊断；新模型 F11_C 从同一公开底座重新训练，保持 140 步预算。',
'''目的：用最小的数据干预检验父职责覆盖不足这一解释。
原缺口：A4 的单 child fused 示范把 prepare 与 aggregate 一起委托，父节点主要做解析与格式化，未覆盖只委托 prepare、父节点承担 aggregate 的状态。
干预：在既有 50 个 TRAIN root 上新执行合法重分配示范，所有可见与 TRAIN 私有检查通过。导出 450 条合法动作，但只把其中 50 条 parent write_code 正例加入修订池。参考实现来自 TRAIN，未用开发题参考代码构造训练目标。
匹配：在原 F11 2,240 个采样槽位的 165 个父代码槽位中，按每 root 隔一个替换一个的规则替换 96 个，覆盖全部 50 个 root；其余 2,144 个槽位保持原样。种子、计划/组件/consume/control 示例、训练超参数及 evaluator 均保持。
实际剂量：数据池 2,240→2,290，不同曝光样本 953→970；目标 token 258,254→260,403，约 +0.83%。这是数据覆盖的单一主要修订，不是每个 token 完全相同的实验。''',
        ['reports/A6.md','runs/a6/data_comparison.json','configs/a6.yaml'],
        '唯一主要变化：父组合示范覆盖；不追加训练步数，不挑中间 checkpoint。')
D.rect(42,149,878,99,WHITE,radius=6)
D.text(60,164,840,31,'原 F11 采样：2,240 个槽位',22,NAVY,True)
D.rect(61,213,800,16,PALE)
D.rect(61,213,800*96/2240,16,CORAL)
D.text(42,267,878,31,'96 个父代码槽位替换    +    2,144 个槽位保持原样',24,TEAL,True,align='center')
D.table(42,323,[292,293,293],['新增覆盖','共同设置','复评设置'],[
    ['50 个既有 TRAIN root','同底座 / seed 42 / 140 步','同 8 root × 5 条件'],
    ['父节点自行 aggregate','计划 / 组件 / 控制不变','同 5 个自由任务'],
    ['仅增加 50 条父代码正例','目标 token +0.83%','主指标：父职责重分配'],
],rowh=34,size=13)

# 18 — A6 result
D.start('17 / A6 · 结果','父职责重分配增加 1 个成功，同时存在回归',
        '同一开发协议配对比较 F11 与 F11_C；图中保留不同条件的实际分子 / 分母。',
'''目的：在原封不动的 A5 evaluator 上检验一次覆盖修订的效果，主要指标预先固定为单 child、parent 保留汇总的通过数。
结果：主指标 1/8→2/8，增加 absolute_ranks，保留 digit_sums。固定组件 12/16→14/16；局部计划+E_ref 3/8→5/8；双 child 组合 7/8→6/8；自由运行 1/5→1/5。
解释：新增父职责示范改变了部分代码行为并带来一个新成功，支持有限覆盖收益。局部计划改善伴随原先 schema 注册失败消失，8 个计划仍全部 ATOMIC；不能声称分解能力提升。
回归：双 child 中 sorted_squares 新成功，但 reverse_negate 和 absolute_ranks 因消费状态或 alias 错误无父程序，净少一个成功。自由运行成功任务也从词序去重换成括号配对，三个 APPS 仍 0/3，不能用总数相同掩盖样本变化。
结论：保留正向变化和回归，一次修订按预定规则结束；F11_C 不是所有指标更好的新最优模型。''',
        ['reports/A6.md','runs/a6/result.json'],
        '主指标 +12.5 pp 只对应新增 1/8 成功；局部改善与退步同时报告。')
D.image(42,146,878,304,charts['a6'])

# 19 — A6 cases
D.start('18 / A6 · 结果解释','源码案例说明：开始承担职责，仍需完整履约',
        '案例来自配对评测和源码诊断；下面是行为概括，不是补写的模型答案。',
'''目的：解释分数变化背后模型究竟做了什么，避免只报告平均值。
新增成功：absolute_ranks 中，F11_C 使用 transform 返回的有序去重列表，计算加权绝对值之和 sum(i * abs(v) ...)，并按根契约直接返回 JSON 数值。它完成了先前欠缺的父汇总职责。
仍失败：sorted_median 和 sorted_squares 已写出所需汇总计算，却返回带 value 或 sum 字段的对象，当前契约要求直接 JSON 结果。因此即便算法步骤有所改善，最终仍失败；这属于输出契约不符，不是无关紧要的空格问题。
回归：双 child 的 reverse_negate 和 absolute_ranks 遇到消费状态/alias 错误，没能产生父程序。算法实现能力和协议稳定性是不同问题。
结论：完整履约包含正确算法、精确输出契约和稳定使用依赖。该组案例支持有限行为变化；没有把单 child 功能通过自动提升成所有责任语义已证明。''',
        ['reports/A6.md','runs/a6/evaluation/F11_C/result.json'] if (ROOT/'runs/a6/evaluation/F11_C/result.json').exists() else ['reports/A6.md','runs/a6/result.json'],
        '完整履约 = 正确计算 + 精确输出契约 + 稳定消费；缺少任何一项都会失败。')
D.card(42,147,282,303,'新增成功 / absolute_ranks','实际消费并完成汇总','使用 child 的有序去重列表\n计算父节点负责的加权和\n直接返回 JSON 数值\n\n主指标新增 1 个成功')
D.card(340,147,282,303,'仍失败 / median、squares','会计算，输出契约仍错','已实现汇总计算\n却返回带 value / sum 的对象\n当前契约要求直接结果\n\n功能判定仍为失败',CORAL)
D.card(638,147,282,303,'组合回归 / 两个任务','消费状态与 alias 出错','reverse_negate\nabsolute_ranks\n未生成有效父程序\n\n双 child 净减少 1 个成功',CORAL)

# 20 — conclusions and A7
D.start('19 / 阶段结论与下一步','下一步补齐“责任结构 → 修复价值”的证据',
        'A7 尚未执行：固定失败集和固定修复器两轨，共用预算与检查权限。',
'''总结：框架闭环和可追踪训练数据已完成；A5 清楚显示固定契约履约与双子模块组合收益；职责重分配暴露了更深的能力缺口，A6 一次干预带来有限改善。A2 还提供了责任筛选支持受控修复的初步依据。
A7 轨道一固定失败源码、证据、修复器和预算，比较定位和修改边界方法，包含受控故障与自然失败。只有有可靠标签的注入故障报告定位命中，自然失败不制造唯一根因标签。
轨道二比较 F00/F10/F01/F11 生成的结构，共用同一定位器与修复器；受控注入尽量对应同一外部义务，并报告全 root 的预算内成功。不能只比较各模型自己的失败修复率，因为失败难度可能不同。
最小对照：无修复、整程序修复/重试、仅提示局部区域、责任证据+ticket+冻结补集+可见重放。指标包括私有成功、边界大小、越界、回归和总成本。
当前停在一次修订之后，继续既定 A7；DPO/RL 不是前置条件。正式封存评测与多 seed 稳定性尚未完成，今天的结果属于开发性证据。''',
        ['STATUS.md','reports/A2.md','reports/A6.md'],
        '核心待检验：在相同修复器与预算下，责任信息和生成期结构是否让修复更可靠？')
D.card(42,149,430,224,'A7 / 轨道一','固定失败，比较修复边界','相同失败源码与可见证据\n相同补丁模型与调用 / token / 测试预算\n衡量修复成功、修改范围与回归')
D.card(490,149,430,224,'A7 / 轨道二','固定修复器，比较生成结构','F00 / F10 / F01 / F11 生成不同结构\n共用定位器与修复器\n同时报告全 root 预算内成功')
D.rect(42,391,878,58,PALE,radius=6)
D.text(58,403,846,43,'最小对照：无修复  |  整程序修复  |  提示局部区域  |  责任证据 + 强制冻结',17,NAVY,True)

# 21 — complete A5 matrix
D.start('备份 01 / 完整结果','A5 四组结果：保留全部条件与分母',
        '每个隔离条件 8 个开发 root；组件合计 16 个条件来自相同 8 root。',
'''用途：导师追问某组或某条件时直接展示完整结果。
本表来自 runs/a5/result.json 与 reports/A5.md。四组分别为功能对照、仅 D、仅 I、D+I。每个单条件有 8 个 root；自由运行有 5 个 root。所有无程序、协议失败、弃权均保留在分母。
双子模块“功能+责任证据”使用保守源码规则：parent 只解析输入、调用已选择 transform、把其返回值传给已选择 summarize、格式化结果，且没有额外算法或别名覆盖。其他写法保留 UNKNOWN，不能把 UNKNOWN 算成责任违规。
各组首个与最终自由成功相同，是因为没有修复尝试。不要把固定组件合计、两种父组合合计当作新的独立任务集合，也不要拿同一 root 的多个条件做未经设计的独立样本显著性检验。''',
        ['reports/A5.md','runs/a5/result.json'])
rows=[
    ['局部计划 + E_ref',*[f'{v}/8' for v in a5values('local_plan')]],
    ['固定 transform',*[f'{v}/8' for v in a5values('transform')]],
    ['固定 summarize',*[f'{v}/8' for v in a5values('summarize')]],
    ['固定组件合计',*[f'{A5["rates"][g]["implementation"]["successes"]}/16' for g in GROUPS]],
    ['双子模块组合',*[f'{v}/8' for v in a5values('parent_both_children')]],
    ['单 child / 父节点汇总',*[f'{v}/8' for v in a5values('parent_transform_child')]],
    ['双 child 功能 + 责任证据','0/8','0/8','8/8','7/8'],
    ['自由：首个 / 最终成功','1/5 · 1/5','1/5 · 1/5','2/5 · 2/5','1/5 · 1/5'],
]
D.table(42,146,[310,142,142,142,142],['评测条件',*GROUPS],rows,rowh=37,size=14)

# 22 — training matching
D.start('备份 02 / 训练剂量与可比性','匹配更新预算；如实报告 token 与重复曝光差异',
        '同一 Qwen3-8B、seed 42、50 个 TRAIN root；每组 2,240 次曝光、140 步更新。',
'''用途：回答四组是否公平、训练是否真实完成，以及为什么不根据 loss 选模型。
匹配项：固定公开底座 revision，NF4/BF16、all-linear LoRA r16/alpha32/dropout0.05、43,646,976 个训练参数、8,192 长度、batch1×累积16、学习率 1e-4、cosine、3% LR warmup；共同采样槽位与 seed42。前 20% 更新使用履约/控制课程。
不匹配项：数据池去重大小、实际曝光的不同事件、目标 token 与输入 token。F00/F10/F01/F11 的不同曝光数为 404/646/711/953。A6 为 970。因监督内容改变，这些并不是仅标签变化或 token 严格等量的消融。
完成证据：五组均完成 140 步，LoRA 参数真实变化，梯度检查与磁盘重载、运行回接通过；固定输入重载 logits 最大差为 0。训练 loss 因目标不同不能用于组间能力排名。共享服务器墙钟受负载影响，不用训练时长比较速度。
A6 保持相同初始参数和槽位顺序，仅替换 96 个父组合槽位；目标 token 约增加 0.83%。''',
        ['reports/A5.md','reports/A6.md','runs/a5/data_comparison.json'])
D.table(42,148,[90,150,169,159,175,135],['模型','去重数据池','不同曝光事件','目标 token','输入 token','更新次数'],[
    ['F00','1,440','404','277,692','5,621,689','140'],
    ['F10','1,850','646','313,447','5,614,258','140'],
    ['F01','1,830','711','222,499','5,913,849','140'],
    ['F11','2,240','953','258,254','5,906,418','140'],
    ['F11_C','2,290','970','260,403','5,893,919','140'],
],rowh=44,size=14)
D.text(42,431,878,60,'所有正式组均完成权重保存、磁盘重载和运行回接；重载前后最大 logits 差异为 0。',17,TEAL,True)

# 23 — metrics
D.start('备份 03 / 指标解释','每个分数在测什么，哪些结论不能直接推导',
        '所有指标依照真实运行与证据强度解释。',
'''FirstProgramSuccess：第一次完整根程序及其依赖闭包的私有通过率，分母为全部选中 root。生成根程序之前可能已做子节点验证和修复，因此不是一次模型调用的 Pass@1。
BudgetedSuccess：终止时当前根产物的私有通过率，即使预算耗尽或弃权，已有产物仍被评分，终止原因另报。首个和最终程序分别冻结，后续修复不能改写首个快照。
功能与责任：功能 PASS 是通过适用测试；责任支持需要明确的契约、源码和依赖证据。UNKNOWN 表示证据不足，不是违约，也不能当责任正例。没有适用局部 oracle 时不能继承根 PASS。
成本：模型调用、输入/输出 token、在线验证与重放、离线评分单独记录；失败补丁成本保留。无成功时每成功成本为 N/A。
数据范围：A2 的20题来自同一 family；A5/A6 隔离用其中8题，自由用5题。私有测试不参与在线选择，但开发结果已被用于诊断和一次修订，所以不能称最终未见泛化测试。''',
        ['reports/F9.md','docs/data_check_protocol_v1.md','reports/A5.md'])
D.table(42,146,[235,385,258],['指标 / 记录','正确解释','需要避免的推断'],[
    ['FirstProgramSuccess','首个完整根程序的私有通过','不等于单次调用 Pass@1'],
    ['BudgetedSuccess','终止时当前产物的私有通过','不遗漏弃权或预算失败'],
    ['功能 PASS','通过适用的功能测试','不自动等于责任正确'],
    ['责任 UNKNOWN','没有足够证据确定责任状态','不当正例，也不当违约'],
    ['完整分母','保留无程序、协议失败、弃权','不只统计已完成程序'],
    ['单 seed 开发差','当前设置下的观察结果','不声称显著性或最终泛化'],
    ['成本','模型与在线 / 离线检查分别记','不把验证减少当全部降本'],
],rowh=42,size=14)

# 24 — infrastructure smoke
D.start('备份 04 / 框架实验与验收','早期 smoke 验证了执行闭环，结果不作为训练收益',
        'F1–F9 与 A3 已完成；下表区分脚本验收、权重连通与真实模型结果。',
'''用途：回答递归、修复、训练与私有隔离是否真正运行过。
F1/F2 完成原子动作状态机及隔离验证，FX01 六个用例通过；F3 三层递归 fixture 为4节点、最大深度2，四节点均功能通过，19次脚本策略动作，责任仍保留 UNKNOWN。
F4 两个 fixture 分别完成只修根代码和计划子树重建，均从失败变为通过；前者冻结 child，后者保持 incoming contract 并重建、重验、重新消费子树。这是执行机制验收，不是真实模型已经学会修复的证据。
F5 接入真实 Qwen3-8B，F6 导出32条一步动作，F7 进行10步 QLoRA smoke、参数变化和重载最大差0；F8 开关 adapter logits 最大差28.625、恢复差0，证明权重生效。
F9 统一评价：脚本6个fixture首个3/6、最终5/6；base五个开发题2/5，smoke adapter为0/5。0/5 是保留的负结果，连通通过不等于能力提升。脚本与模型任务不同，不能跨组公平比较；F9 与 A5 预算和训练设置也不同。
A3 冻结协议后六个fixture最终5/6，其中一个为预设错误父程序；最新A6全量261项测试通过。各阶段测试数是累计套件规模，不能逐阶段相加。''',
        ['reports/F3.md','reports/F4.md','reports/F7.md','reports/F9.md','reports/A3.md'])
D.table(42,147,[120,415,343],['阶段','实际验证','解释'],[
    ['F1–F2','原子生成 → 隔离验证；FX01 6 例通过','原子执行闭环'],
    ['F3','三层递归；4 个节点均功能 PASS','显式消费与父组合可执行'],
    ['F4','代码修复 / 计划重建均 FAIL → PASS','脚本驱动，验证边界与回滚'],
    ['F5–F8','真实模型、32条导出、10步训练、回接','证明新权重确实参与推理'],
    ['F9','base 2/5；smoke adapter 0/5','连通验收含负结果，非训练收益'],
    ['A3','冻结主机制；fixture 最终 5/6','一个预设错误 parent 保留失败'],
    ['最近回归','A6：261 项测试通过','代码检查通过，不代替能力结果'],
],rowh=42,size=13.5)

# 25 — FAQ
D.start('备份 05 / 常见追问','当前结论与实验边界可以这样回答',
        '主证据索引与详细解释均已写入每页演讲备注及配套讲解稿。',
'''问：为什么层次契约在 A2 成功率低，还继续这条路线？
答：A2 没证明零样本生成更强，但揭示了 child 已通过而 parent 仍失败的明确缺口；契约提供可绑定局部检查的学习对象。后续 A5 观察到固定契约履约和双 child 组合的训练响应，支持继续检验这条方向。
问：最强的阶段结果是什么？
答：F01 相对 F00 的固定组件 0/16→13/16、双 child 组合1/8→8/8；同时职责重分配0/8暴露了常见组合之外的能力缺口。前者是正向证据，后者是具体研究问题。
问：F11_C 是否已经全面优于 F11？
答：没有。主指标1/8→2/8、组件12/16→14/16，但双 child7/8→6/8，自由成功总数不变且通过任务替换，因此报告有限改善和回归。
问：论文核心还缺什么证据？
答：生成期责任结构对固定修复器的下游价值仍缺 A7 的决定性比较；多 seed 稳定性及正式封存评测还未完成。下一步不扩大工程范围，按两轨实验补这个环节。
问：是否有数据泄漏或挑选 checkpoint？
答：TRAIN 与开发按 root/family 隔离，私有评分不参与在线选择，五组都取140步最终权重。A5/A6 使用已消费开发集，修订也基于开发诊断，因此如实称开发比较，不能称最终未见测试。''',
        ['reports/A2.md','reports/A4.md','reports/A5.md','reports/A6.md'])
items=[('为什么仍保留层次契约？','它提供可检查的责任学习对象；A5 已看到固定契约训练响应。'),
       ('F11_C 是全面升级吗？','主指标和组件改善，双 child 回归；结论是有限收益。'),
       ('分解与修复学会了吗？','主动 SPLIT 和训练后结构的修复收益尚无完成的证据。'),
       ('下一步最关键的实验？','相同失败 / 修复器 / 预算下，检验责任边界与生成结构价值。')]
for i,(q,a) in enumerate(items):
    yy=146+i*80
    D.rect(42,yy,878,70,WHITE if i%2==0 else PALE,radius=5)
    D.text(58,yy+10,842,27,q,19,TEAL,True)
    D.text(58,yy+40,842,27,a,16,INK)

assert len(D.slides)==25
D.save()

# Preview renders are generated from the same positioned elements as the PPTX.
doc=fitz.open(OUT/f'{STEM}.pdf')
thumbs=[]
for i,p in enumerate(doc):
    pix=p.get_pixmap(matrix=fitz.Matrix(1.25,1.25),alpha=False)
    path=ASSETS/f'slide_{i+1:02d}.png'
    pix.save(path)
    im=Image.open(path).convert('RGB')
    im.thumbnail((384,216))
    thumbs.append(im.copy())
sheet=Image.new('RGB',(5*400,5*242),'#DCE5EA')
draw=ImageDraw.Draw(sheet)
for i,im in enumerate(thumbs):
    xx=(i%5)*400+8; yy=(i//5)*242+5
    sheet.paste(im,(xx,yy))
    draw.text((xx+5,yy+218),f'{i+1:02d}',fill='#17364B')
sheet.save(OUT/'全部页面预览.png')
selected=Image.new('RGB',(1200,676),'#DCE5EA')
for j,i in enumerate([3,12,13,17]):
    im=Image.open(ASSETS/f'slide_{i+1:02d}.png').convert('RGB')
    im.thumbnail((592,333))
    selected.paste(im,((j%2)*600+4,(j//2)*338+3))
selected.save(OUT/'核心页面预览.png')

# Focused artifact checks: page count, notes, evidence links, saved result agreement.
check=Presentation(OUT/f'{STEM}.pptx')
assert len(check.slides)==len(doc)==25
assert all(len(s.notes_slide.notes_text_frame.text)>100 for s in check.slides)
missing=[p for s in D.slides for p in s['source'] if not (ROOT/p).exists()]
assert not missing,missing
assert [A5['rates'][g]['implementation']['successes'] for g in GROUPS]==[0,3,13,12]
assert a5values('parent_both_children')==[1,1,8,7]
assert a5values('parent_transform_child')==[0,0,0,1]
assert A6['primary_metric']['baseline_passes']==1 and A6['primary_metric']['revised_passes']==2
print(json.dumps({'slides':len(check.slides),'pdf_pages':len(doc),'notes':'25/25',
                  'pptx':str(OUT/f'{STEM}.pptx'),'pdf':str(OUT/f'{STEM}.pdf'),
                  'handout':str(OUT/f'{STEM}_逐页讲解稿.md'),'evidence_links':'all exist'},ensure_ascii=False,indent=2))
